# -*- coding: utf-8 -*-
"""
在原供应商 Excel 上生成「计算版」：

1. 完整复制工作簿；
2. 明细行：与资产管理表按「目录+类型+名称+开通/关停+数量+单价」匹配，写回核算/应计费/合计等列；
3. 「本月合计费用」汇总区：按每 Sheet 公式后的明细（与界面「各 Sheet」一致）按「服务目录编号+服务名称」聚合，
   将合计费用、本月核算金额、数量、服务类型写回汇总表（表头无「开通时间」、有「合计费用」），并写「总计金额」行；
4. 工作簿第一张表上的总览/总账统计表：优先按「资产管理」全量表（生成计算版时传入的 detail_df，与 数据库资产.csv / 程序计算明细 同源）
   按「服务目录编号+服务名称」聚合合计费用、本月核算金额后回填（服务名称经 _ledger_agg_service_name_key 归一，简写与「类名-英文产品名」合并）；仅当全表无法聚合时才回退为各 Sheet 内存试算结果；
5. 追加「程序计算明细」子表。

需要 openpyxl。
"""
from __future__ import annotations

import os
import re
import shutil
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from data_manager import (
    BILLING_DAYS_COLUMN,
    parse_money_decimal,
    round_money_half_up,
    sum_parsed_money,
)
from sheet_normalize import (
    SheetMatchRecord,
    _VENDOR_BILL_DAYS_HEADER_RE,
    _normalize_directory_id_token,
    _row_as_list,
    _strip_cell,
    extract_sheet_match_records,
    is_empty_row,
    is_header_row,
    is_title_row,
    parse_data_row,
    try_get_ledger_cost_table_header_idx,
    try_get_ledger_month_table_header_idx,
    try_get_ledger_unified_overview_header_idx,
    try_get_month_fee_summary_header_idx,
)

DEFAULT_DETAIL_SHEET = "程序计算明细"

_WRITE_HEADER_ALIASES: Dict[str, Tuple[str, ...]] = {
    "本月核算天数": ("本月核算天数",),
    "本月核算金额": ("本月核算金额",),
    BILLING_DAYS_COLUMN: (BILLING_DAYS_COLUMN, "截至2026年4月30日应计费天数"),
    "合计费用": ("合计费用", "本期合计费用"),
    "数量": ("数量",),
    "服务类型": ("服务类型",),
}


def _sanitize_excel_sheet_title(name: str) -> str:
    s = (name or "").strip() or DEFAULT_DETAIL_SHEET
    for ch in "[]\\:*?/":
        s = s.replace(ch, "_")
    s = re.sub(r"\s+", " ", s).strip()
    return s[:31]


def _date_key(val: Any) -> str:
    s = str(val or "").strip()
    if not s:
        return ""
    ts = pd.to_datetime(s, errors="coerce")
    if pd.isna(ts):
        return s[:19].replace("/", "-") if s else ""
    return str(ts.date())


def _qty_key(val: Any) -> str:
    s = str(val or "").strip().replace(",", "").replace("，", "")
    if not s:
        return ""
    try:
        x = float(s)
        if abs(x - round(x)) < 1e-9:
            return str(int(round(x)))
        return f"{x:.6g}"
    except ValueError:
        return s


def _price_key(val: Any) -> str:
    s = str(val or "").strip().replace("￥", "").replace("¥", "").replace(",", "").replace("，", "")
    if not s:
        return ""
    try:
        x = float(s)
        return f"{x:.6g}"
    except ValueError:
        return s


def _norm_text(val: Any) -> str:
    return re.sub(r"\s+", " ", str(val or "").strip())


def _ledger_agg_service_name_key(name: Any) -> str:
    """总览/各 Sheet 按「目录编号+服务名称」聚合时使用的名称键。

    供应商常把同一产品写成简写与「中文类名-品牌」两行（如 OceanBase 与 国产关系型数据库-OceanBase），
    若仅按整串匹配会拆成两条导致首页或汇总金额偏少。规则：名称中含连字符时，若最后一段含连续拉丁字母
    （典型英文产品名）则用该段作为键；尾段至少 3 个拉丁字母才启用，避免「AB」类误把不同产品并成一桶（程序侧合计偏高）。
    否则用整行名称；再统一 casefold 比较。
    """
    s = _norm_text(name)
    if not s:
        return ""
    for ch in "－—–−":
        s = s.replace(ch, "-")
    if "-" in s:
        tail = s.rsplit("-", 1)[-1].strip()
        if len(tail) >= 2 and re.search(r"[A-Za-z]{3,}", tail):
            s = tail
    return s.casefold()


def _match_key_from_parsed(p: Dict[str, str]) -> Tuple[str, str, str, str, str, str, str]:
    return (
        _normalize_directory_id_token(p.get("服务目录编号", "")),
        _norm_text(p.get("服务类型", "")),
        _norm_text(p.get("服务名称", "")),
        _date_key(p.get("开通时间", "")),
        _date_key(p.get("关停时间", "")),
        _qty_key(p.get("数量", "")),
        _price_key(p.get("单价", "")),
    )


def _match_key_from_series(row: pd.Series) -> Tuple[str, str, str, str, str, str, str]:
    return (
        _normalize_directory_id_token(row.get("服务目录编号", "")),
        _norm_text(row.get("服务类型", "")),
        _norm_text(row.get("服务名称", "")),
        _date_key(row.get("开通时间", "")),
        _date_key(row.get("关停时间", "")),
        _qty_key(row.get("数量", "")),
        _price_key(row.get("单价", "")),
    )


def _build_asset_lookup(assets_df: pd.DataFrame) -> Tuple[Dict[Tuple[str, ...], pd.Series], int]:
    out: Dict[Tuple[str, str, str, str, str, str, str], pd.Series] = {}
    dup = 0
    for _, row in assets_df.iterrows():
        k = _match_key_from_series(row)
        if k in out:
            dup += 1
            continue
        out[k] = row
    return out, dup


def _header_col_j0(header_idx: Dict[str, int], logical: str) -> Optional[int]:
    for name in _WRITE_HEADER_ALIASES.get(logical, (logical,)):
        if name in header_idx:
            return header_idx[name]
    if logical == BILLING_DAYS_COLUMN:
        for k, j in header_idx.items():
            if _VENDOR_BILL_DAYS_HEADER_RE.fullmatch(k):
                return j
    return None


def _excel_cell_value(logical: str, raw: Any) -> Any:
    if raw is None:
        return None
    if isinstance(raw, float) and pd.isna(raw):
        return None
    st = str(raw).strip()
    if logical in ("本月核算天数", BILLING_DAYS_COLUMN):
        if not st:
            return 0
        try:
            return max(0, int(float(st.replace(",", ""))))
        except ValueError:
            return st
    s2 = st.replace("￥", "").replace("¥", "").replace(",", "").replace("，", "")
    if not s2:
        return 0.0
    if logical in ("本月核算金额", "合计费用"):
        try:
            return round_money_half_up(parse_money_decimal(s2))
        except ValueError:
            return st
    try:
        return float(s2)
    except ValueError:
        return st


def _write_money_cell(ws, row_1based: int, col_1based: int, val: Any) -> None:
    """写入金额单元格：数值保留到分，并设 Excel 数字格式为两位小数，避免显示成整数或丢尾数。"""
    v = round_money_half_up(val)
    cell = ws.cell(row=row_1based, column=col_1based)
    cell.value = v
    cell.number_format = "#,##0.00"


def _aggregate_sheet_for_month_summary(rc: pd.DataFrame) -> Dict[Tuple[str, str], Dict[str, Any]]:
    """按（服务目录编号, 服务名称）聚合本 Sheet 公式后的金额与数量。"""
    if rc is None or len(rc) == 0:
        return {}
    need = ("服务目录编号", "服务名称", "合计费用", "本月核算金额", "数量")
    for c in need:
        if c not in rc.columns:
            return {}
    d = rc.fillna("")
    d = d.copy()
    d["_c"] = d["服务目录编号"].astype(str).map(_normalize_directory_id_token)
    d["_n"] = d["服务名称"].astype(str).map(_ledger_agg_service_name_key)
    out: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for (c, n), g in d.groupby(["_c", "_n"], sort=False):
        if not c.isdigit():
            continue
        sum_cost = sum_parsed_money(g["合计费用"])
        sum_month = sum_parsed_money(g["本月核算金额"])
        sum_qty = float(
            pd.to_numeric(
                g["数量"].astype(str).str.replace(",", "").str.replace("，", ""),
                errors="coerce",
            )
            .fillna(0)
            .sum()
        )
        st = ""
        if "服务类型" in g.columns:
            st = _norm_text(g["服务类型"].iloc[0])
        out[(c, n)] = {
            "合计费用": sum_cost,
            "本月核算金额": sum_month,
            "数量": sum_qty,
            "服务类型": st,
        }
    return out


def _merge_workbook_ledger_aggregates(
    per_sheet_calc: List[Tuple[str, pd.DataFrame]],
    first_sheet_name: str,
    detail_sheet_name: str,
) -> Dict[Tuple[str, str], Dict[str, Any]]:
    """除首页、明细子表外，各 Sheet 按（目录编号, 服务名称）聚合金额后跨表相加（统计学汇总，不重算单价×天）。"""
    merged: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for sn, rc in per_sheet_calc:
        if sn == first_sheet_name or sn == detail_sheet_name:
            continue
        agg = _aggregate_sheet_for_month_summary(rc)
        for key, a in agg.items():
            e = merged.setdefault(
                key,
                {
                    "合计费用": Decimal("0"),
                    "本月核算金额": Decimal("0"),
                    "数量": 0.0,
                    "服务类型": str(a.get("服务类型") or ""),
                },
            )
            e["合计费用"] += parse_money_decimal(a["合计费用"])
            e["本月核算金额"] += parse_money_decimal(a["本月核算金额"])
            e["数量"] += float(a["数量"])
            if not e.get("服务类型") and a.get("服务类型"):
                e["服务类型"] = str(a["服务类型"])
    for e in merged.values():
        e["合计费用"] = round_money_half_up(e["合计费用"])
        e["本月核算金额"] = round_money_half_up(e["本月核算金额"])
    return merged


def _lookup_ledger_row_aggregate(
    merged: Dict[Tuple[str, str], Dict[str, Any]],
    cid_raw: Any,
    name_raw: Any,
) -> Optional[Dict[str, Any]]:
    """首页总览一行对应聚合结果：先按（目录编号, 服务名称）精确匹配。

    模板上的目录编号可能与库不一致时，按「服务名称」在聚合结果里查找；若同名对应多条
    （不同目录编号），**不得**把多条金额相加（易把「人大金仓」等加成约 3 倍），应：
    优先取与首页行目录编号一致的一条；若无，则取目录编号数值与首页最接近的一条（如 EulerOS）。
    距离相同时取合计费用较小者，减轻误选邻近高价桶（程序侧比供应商表偏高）。
    若首页编号解析为空/非数字：按目录编号升序，合计费用较小者优先，不用 dict 迭代顺序的 hits[0]。
    """
    c = _normalize_directory_id_token(str(cid_raw or ""))
    n = _ledger_agg_service_name_key(str(name_raw or ""))
    if not n:
        return None
    key = (c, n)
    if key in merged:
        return merged[key]
    hits: List[Tuple[str, Dict[str, Any]]] = [
        (k[0], v) for k, v in merged.items() if k[1] == n and k[0].isdigit()
    ]
    if not hits:
        return None
    if len(hits) == 1:
        return hits[0][1]
    same_c = [v for cid, v in hits if cid == c]
    if len(same_c) == 1:
        return same_c[0]
    if len(same_c) > 1:
        st = ""
        for v in same_c:
            t = str(v.get("服务类型") or "").strip()
            if t and not st:
                st = t
        return {
            "合计费用": sum_parsed_money(x["合计费用"] for x in same_c),
            "本月核算金额": sum_parsed_money(x["本月核算金额"] for x in same_c),
            "数量": sum(float(x["数量"]) for x in same_c),
            "服务类型": st or _norm_text(str(same_c[0].get("服务类型") or "")),
        }
    if c.isdigit():
        try:
            target = int(c)

            def _cid_sort_key(item: Tuple[str, Dict[str, Any]]) -> Tuple[int, Decimal]:
                try:
                    cid = int(item[0])
                    dist = abs(cid - target)
                    cost = parse_money_decimal(item[1].get("合计费用", 0))
                    return (dist, cost)
                except (ValueError, TypeError):
                    return (Decimal(10**9), Decimal(10**9))

            return min(hits, key=_cid_sort_key)[1]
        except ValueError:
            pass
    # 编号为空或非数字：按目录编号升序，合计费用较小者优先（避免误选高价桶）
    try:
        return min(
            hits,
            key=lambda item: (int(item[0]), parse_money_decimal(item[1].get("合计费用", 0))),
        )[1]
    except (ValueError, TypeError):
        return hits[0][1]


def _sliding_find_ledger_header_row(row_vals: List[str], kind: str) -> Optional[Dict[str, int]]:
    """在整行表头中滑动窗口，匹配左/右并列总览表。窗口不宜过宽，否则会同时框进两侧表头导致无法识别。

    左侧「合计费用」表：优先最靠左的匹配。右侧「本月核算金额」表：不得先用「统一六列」匹配到左侧
    （统一识别器只认「合计费用」列名，会误把左表当成右表），应先试本月核算表头，必要时再在靠右
    位置用统一六列兜底（两表最后一列都叫「合计费用」的模板）。
    """
    n = len(row_vals)
    best: Optional[Dict[str, int]] = None
    best_key: Optional[Tuple[int, int]] = None  # cost: 最小 (c0, win_w)；month: 最大 c0，同列起点取更窄 win_w
    for win_w in (10, 12, 14, 16, 18):
        for c0 in range(0, min(max(n, 1), 96)):
            win = [
                _strip_cell(row_vals[j]) if j < len(row_vals) else ""
                for j in range(c0, min(c0 + win_w, n))
            ]
            if len(win) < 4:
                continue
            if kind == "cost":
                hi_rel = try_get_ledger_unified_overview_header_idx(win) or try_get_ledger_cost_table_header_idx(
                    win
                )
            else:
                hi_rel = try_get_ledger_month_table_header_idx(win)
                if not hi_rel:
                    hi_u = try_get_ledger_unified_overview_header_idx(win)
                    # 仅当窗口起点已越过左侧表（约第 9 列起），才用「六列同名」兜底，避免误匹配左表
                    if hi_u and c0 >= 8:
                        hi_rel = hi_u
            if not hi_rel:
                continue
            key = (c0, win_w)
            if kind == "month":
                if best_key is None or key[0] > best_key[0] or (
                    key[0] == best_key[0] and key[1] < best_key[1]
                ):
                    best_key = key
                    best = {k: c0 + int(v) for k, v in hi_rel.items()}
            else:
                if best_key is None or key < best_key:
                    best_key = key
                    best = {k: c0 + int(v) for k, v in hi_rel.items()}
    return best


def _first_sheet_ledger_title_hit(row_t: List[str], markers: Tuple[str, ...]) -> bool:
    joined = "".join(_strip_cell(x) for x in row_t)
    return any(m in joined for m in markers)


def _row_has_ledger_grand_total_label(vals: List[str]) -> bool:
    s = "".join(_strip_cell(v) for v in vals[:18])
    return "费用总合计" in s


def _scan_first_ledger_table_by_title(
    dfg: pd.DataFrame,
    title_markers: Tuple[str, ...],
    kind: str,
) -> Optional[Tuple[int, Dict[str, int], List[Tuple[int, Dict[str, str]]], Optional[int]]]:
    """标题行命中任一 title_markers 下的总览/总账表：返回 (表头 Excel 行, 列索引, 数据行, 费用总合计行)。"""
    n = len(dfg)
    base_cols = int(dfg.shape[1]) if dfg.shape[1] else 0
    max_cols = max(base_cols, 28)
    t = 0
    while t < n:
        row_t = _row_as_list(dfg.iloc[t], max_cols + 40)
        if not _first_sheet_ledger_title_hit(row_t, title_markers):
            t += 1
            continue
        hi_abs: Optional[Dict[str, int]] = None
        h_idx = -1
        for hh in range(t + 1, min(t + 50, n)):
            vals = _row_as_list(dfg.iloc[hh], max_cols + 40)
            hi_abs = _sliding_find_ledger_header_row(vals, kind)
            if hi_abs:
                h_idx = hh
                break
        if not hi_abs:
            t += 1
            continue
        k = h_idx + 1
        data_targets: List[Tuple[int, Dict[str, str]]] = []
        total_excel: Optional[int] = None
        while k < n:
            vals2 = _row_as_list(dfg.iloc[k], max_cols + 40)
            if _row_has_ledger_grand_total_label(vals2):
                total_excel = k + 1
                break
            if is_header_row(vals2):
                break
            if is_title_row(vals2):
                break
            if is_empty_row(vals2):
                k += 1
                continue
            parsed = parse_data_row(hi_abs, vals2)
            cid = _normalize_directory_id_token(parsed.get("服务目录编号", ""))
            if cid.isdigit():
                data_targets.append((k + 1, parsed))
            k += 1
        return (h_idx + 1, hi_abs, data_targets, total_excel)

    return None


def _scan_month_fee_blocks(
    dfg: pd.DataFrame,
) -> List[Tuple[List[Tuple[int, Dict[str, int], Dict[str, str]]], Optional[int], Dict[str, int]]]:
    """扫描「本月合计费用」标题下的汇总块：[(数据行…, 总计行号或 None, 表头索引), …]。"""
    blocks: List[
        Tuple[List[Tuple[int, Dict[str, int], Dict[str, str]]], Optional[int], Dict[str, int]]
    ] = []
    n = len(dfg)
    max_cols = int(dfg.shape[1]) if dfg.shape[1] else 0
    t = 0
    while t < n:
        vals = _row_as_list(dfg.iloc[t], max_cols)
        if not any("本月合计费用" in _strip_cell(x) for x in vals[:18]):
            t += 1
            continue
        hi: Optional[Dict[str, int]] = None
        h_idx = -1
        for hh in range(t + 1, min(t + 22, n)):
            ncell = len(dfg.iloc[hh])
            hv = [
                _strip_cell(dfg.iloc[hh, j]) if j < ncell else ""
                for j in range(max(ncell, max_cols))
            ]
            hi = try_get_month_fee_summary_header_idx(hv)
            if hi:
                h_idx = hh
                break
        if not hi:
            t += 1
            continue
        k = h_idx + 1
        data_targets: List[Tuple[int, Dict[str, int], Dict[str, str]]] = []
        total_excel: Optional[int] = None
        while k < n:
            vals2 = _row_as_list(dfg.iloc[k], max_cols)
            if is_header_row(vals2) or is_title_row(vals2):
                break
            v0 = _strip_cell(vals2[0]) if vals2 else ""
            if v0 and ("总计金额" in v0 or v0 == "总计金额" or (v0.startswith("总计") and "金额" in v0)):
                total_excel = k + 1
                k += 1
                break
            if is_empty_row(vals2):
                k += 1
                continue
            parsed = parse_data_row(hi, vals2)
            cid = _normalize_directory_id_token(parsed.get("服务目录编号", ""))
            if cid.isdigit():
                data_targets.append((k + 1, hi, parsed))
            k += 1
        blocks.append((data_targets, total_excel, hi))
        t = max(t + 1, k)
    return blocks


def _fill_detail_lines(
    wb,
    dest_xlsx: str,
    assets_df: pd.DataFrame,
    detail_title: str,
) -> Tuple[int, int, int, int]:
    lookup, _dup = _build_asset_lookup(assets_df)
    total_records = 0
    matched_rows = 0
    cells = 0
    unmatched = 0

    xls = pd.ExcelFile(dest_xlsx)
    try:
        sheet_grids: Dict[str, pd.DataFrame] = {}
        for sn in xls.sheet_names:
            if sn == detail_title:
                continue
            sheet_grids[sn] = pd.read_excel(
                dest_xlsx,
                sheet_name=sn,
                header=None,
                dtype=str,
                keep_default_na=False,
            ).fillna("")
    finally:
        xls.close()

    for sn, dfg in sheet_grids.items():
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        records: List[SheetMatchRecord] = extract_sheet_match_records(dfg)
        total_records += len(records)
        for rec in records:
            k = _match_key_from_parsed(rec.parsed)
            row = lookup.get(k)
            if row is None:
                unmatched += 1
                continue
            matched_rows += 1
            for logical in ("本月核算天数", "本月核算金额", BILLING_DAYS_COLUMN, "合计费用"):
                j0 = _header_col_j0(rec.header_idx, logical)
                if j0 is None:
                    continue
                if logical not in row.index:
                    continue
                val = _excel_cell_value(logical, row.get(logical))
                if val is None:
                    continue
                try:
                    c = ws.cell(row=rec.excel_row_1based, column=j0 + 1)
                    c.value = val
                    if logical in ("本月核算金额", "合计费用"):
                        c.number_format = "#,##0.00"
                    cells += 1
                except Exception:
                    pass

    return total_records, matched_rows, cells, unmatched


def _fill_month_fee_summary_sections(
    wb,
    dest_xlsx: str,
    per_sheet_calc: List[Tuple[str, pd.DataFrame]],
    detail_title: str,
) -> str:
    if not per_sheet_calc:
        return ""
    xls = pd.ExcelFile(dest_xlsx)
    try:
        sheet_grids: Dict[str, pd.DataFrame] = {}
        for sn in xls.sheet_names:
            if sn == detail_title:
                continue
            sheet_grids[sn] = pd.read_excel(
                dest_xlsx,
                sheet_name=sn,
                header=None,
                dtype=str,
                keep_default_na=False,
            ).fillna("")
    finally:
        xls.close()

    n_blocks = 0
    n_sum_rows = 0
    n_total_rows = 0
    n_cells = 0

    for sn, rc in per_sheet_calc:
        if sn not in sheet_grids or sn not in wb.sheetnames:
            continue
        agg = _aggregate_sheet_for_month_summary(rc)
        if not agg:
            continue
        dfg = sheet_grids[sn]
        ws = wb[sn]
        for data_targets, total_excel, hi in _scan_month_fee_blocks(dfg):
            if not data_targets and total_excel is None:
                continue
            n_blocks += 1
            block_cost_sum = Decimal("0")
            for exr, hidx, parsed in data_targets:
                key = (
                    _normalize_directory_id_token(parsed.get("服务目录编号", "")),
                    _ledger_agg_service_name_key(parsed.get("服务名称", "")),
                )
                a = agg.get(key)
                if not a:
                    continue
                n_sum_rows += 1
                block_cost_sum += parse_money_decimal(a["合计费用"])
                for logical in ("本月核算金额", "合计费用", "数量"):
                    j0 = _header_col_j0(hidx, logical)
                    if j0 is None:
                        continue
                    try:
                        if logical == "数量":
                            ws.cell(row=exr, column=j0 + 1).value = float(a["数量"])
                        else:
                            _write_money_cell(ws, exr, j0 + 1, a[logical])
                        n_cells += 1
                    except Exception:
                        pass
                jt = _header_col_j0(hidx, "服务类型")
                if jt is not None and a.get("服务类型"):
                    try:
                        ws.cell(row=exr, column=jt + 1).value = a["服务类型"]
                        n_cells += 1
                    except Exception:
                        pass
            if total_excel is not None:
                jc = _header_col_j0(hi, "合计费用")
                if jc is not None:
                    try:
                        _write_money_cell(ws, total_excel, jc + 1, block_cost_sum)
                        n_total_rows += 1
                        n_cells += 1
                    except Exception:
                        pass

    if n_blocks == 0:
        return ""
    return (
        f"\n「本月合计费用」汇总：识别 {n_blocks} 个汇总块，写入 {n_sum_rows} 行产品小计、"
        f"{n_total_rows} 行「总计金额」，共约 {n_cells} 格（按每 Sheet 公式后明细聚合）。"
    )


def _fill_first_sheet_ledger_summaries(
    wb,
    dest_xlsx: str,
    detail_title: str,
    *,
    assets_df: Optional[pd.DataFrame] = None,
    per_sheet_calc: Optional[List[Tuple[str, pd.DataFrame]]] = None,
) -> str:
    """第一张表：总览/总账表（开展/开通 + 本月核算）写合计费用、本月核算金额及费用总合计。"""
    xls = pd.ExcelFile(dest_xlsx)
    try:
        names = list(xls.sheet_names)
        if not names:
            return ""
        first_sn = names[0]
        if first_sn not in wb.sheetnames:
            return ""
        merged: Dict[Tuple[str, str], Dict[str, Any]] = {}
        agg_source = ""
        if assets_df is not None and len(assets_df) > 0:
            ded = assets_df.drop(columns=["数据块"], errors="ignore").copy()
            merged = _aggregate_sheet_for_month_summary(ded) or {}
            if merged:
                agg_source = "资产管理全表"
        if not merged and (per_sheet_calc or []):
            merged = _merge_workbook_ledger_aggregates(
                list(per_sheet_calc or []), first_sn, detail_title
            )
            if merged:
                agg_source = "各Sheet内存试算（回退）"
        if not merged:
            return ""
        dfg = pd.read_excel(
            dest_xlsx,
            sheet_name=first_sn,
            header=None,
            dtype=str,
            keep_default_na=False,
        ).fillna("")
    finally:
        xls.close()

    ws = wb[first_sn]
    n_cells = 0
    n_rows_cost = 0
    n_rows_month = 0
    n_tot = 0

    cost_scan = _scan_first_ledger_table_by_title(
        dfg,
        (
            "服务开展情况总览统计表",
            "服务开通情况总账统计表",
        ),
        "cost",
    )
    if cost_scan:
        _h, hi, data_rows, total_excel = cost_scan
        block_sum = Decimal("0")
        jc = hi.get("合计费用")
        jt = hi.get("服务类型")
        jq = hi.get("数量")
        for exr, parsed in data_rows:
            a = _lookup_ledger_row_aggregate(
                merged, parsed.get("服务目录编号", ""), parsed.get("服务名称", "")
            )
            if not a:
                continue
            n_rows_cost += 1
            if jc is not None:
                v = parse_money_decimal(a["合计费用"])
                _write_money_cell(ws, exr, jc + 1, v)
                block_sum += v
                n_cells += 1
            if jt is not None and a.get("服务类型"):
                try:
                    ws.cell(row=exr, column=jt + 1).value = a["服务类型"]
                    n_cells += 1
                except Exception:
                    pass
            if jq is not None:
                try:
                    ws.cell(row=exr, column=jq + 1).value = float(a["数量"])
                    n_cells += 1
                except Exception:
                    pass
        if total_excel is not None and jc is not None:
            try:
                _write_money_cell(ws, total_excel, jc + 1, block_sum)
                n_tot += 1
                n_cells += 1
            except Exception:
                pass

    month_scan = _scan_first_ledger_table_by_title(
        dfg,
        (
            "本月核算情况总览统计表",
            "本月核算情况总账统计表",
        ),
        "month",
    )
    if month_scan:
        _h, hi, data_rows, total_excel = month_scan
        block_sum_m = Decimal("0")
        # 与左侧总览列名一致时，最后一列在表头中也叫「合计费用」，写入的是本月核算金额汇总
        jm = hi.get("本月核算金额")
        if jm is None:
            jm = hi.get("合计费用")
        jt = hi.get("服务类型")
        jq = hi.get("数量")
        for exr, parsed in data_rows:
            a = _lookup_ledger_row_aggregate(
                merged, parsed.get("服务目录编号", ""), parsed.get("服务名称", "")
            )
            if not a:
                continue
            n_rows_month += 1
            if jm is not None:
                v = parse_money_decimal(a["本月核算金额"])
                _write_money_cell(ws, exr, jm + 1, v)
                block_sum_m += v
                n_cells += 1
            if jt is not None and a.get("服务类型"):
                try:
                    ws.cell(row=exr, column=jt + 1).value = a["服务类型"]
                    n_cells += 1
                except Exception:
                    pass
            if jq is not None:
                try:
                    ws.cell(row=exr, column=jq + 1).value = float(a["数量"])
                    n_cells += 1
                except Exception:
                    pass
        if total_excel is not None and jm is not None:
            try:
                _write_money_cell(ws, total_excel, jm + 1, block_sum_m)
                n_tot += 1
                n_cells += 1
            except Exception:
                pass

    if n_cells == 0:
        return ""
    return (
        f"\n首页总览/总账统计（{agg_source} 聚合；编号不一致时按名称取最近目录编号，不跨编号相加）：合计费用表写入 {n_rows_cost} 行、"
        f"本月核算表写入 {n_rows_month} 行，「费用总合计」{n_tot} 处，共约 {n_cells} 格（未改写使用天数）。"
    )


def export_workbook_copy_with_detail_sheet(
    source_xlsx: str,
    dest_xlsx: str,
    detail_df: pd.DataFrame,
    detail_sheet_name: str = DEFAULT_DETAIL_SHEET,
    *,
    fill_source_columns: bool = True,
    per_sheet_calc: Optional[List[Tuple[str, pd.DataFrame]]] = None,
) -> Tuple[bool, str]:
    if not source_xlsx or not os.path.isfile(source_xlsx):
        return False, "源文件不存在。"
    low = source_xlsx.lower()
    if low.endswith(".xls") and not low.endswith((".xlsx", ".xlsm")):
        return False, "当前仅支持 .xlsx / .xlsm（旧版 .xls 请先另存为 xlsx）。"

    dest_dir = os.path.dirname(os.path.abspath(dest_xlsx))
    if dest_dir and not os.path.isdir(dest_dir):
        os.makedirs(dest_dir, exist_ok=True)

    shutil.copy2(source_xlsx, dest_xlsx)

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.dataframe import dataframe_to_rows
    except ImportError:
        return (
            False,
            f"已复制原表到：\n{dest_xlsx}\n\n"
            "但未继续处理：请先安装 openpyxl。\n"
            "命令行执行：pip install openpyxl",
        )

    title = _sanitize_excel_sheet_title(detail_sheet_name)
    fill_stats = ""
    try:
        wb = load_workbook(dest_xlsx, read_only=False, data_only=False)

        if fill_source_columns and detail_df is not None and len(detail_df) > 0:
            tr, mr, ce, um = _fill_detail_lines(wb, dest_xlsx, detail_df, title)
            fill_stats = (
                f"\n明细自动回填：识别明细行 {tr} 条，匹配成功 {mr} 条，写入单元格 {ce} 个；"
                f"未匹配 {um} 条。\n"
            )
            if per_sheet_calc:
                fill_stats += _fill_month_fee_summary_sections(
                    wb, dest_xlsx, per_sheet_calc, title
                )
            fill_stats += _fill_first_sheet_ledger_summaries(
                wb,
                dest_xlsx,
                title,
                assets_df=detail_df,
                per_sheet_calc=per_sheet_calc,
            )

        if title in wb.sheetnames:
            wb.remove(wb[title])
        ws = wb.create_sheet(title=title)
        df = detail_df.copy() if detail_df is not None else pd.DataFrame()
        df = df.fillna("")
        for row in dataframe_to_rows(df, index=False, header=True):
            ws.append(list(row))
        wb.save(dest_xlsx)
    except Exception as e:
        return False, f"已复制到：{dest_xlsx}\n处理失败：{e}"

    n = len(detail_df) if detail_df is not None else 0
    return (
        True,
        f"已生成计算版：\n{dest_xlsx}\n"
        f"{fill_stats}"
        f"已追加工作表「{title}」（{n} 行，标准列全量）。\n"
        "若未出现「本月合计费用」或首页总览回填：请先完成「清洗、计算并载入」；首页总览金额与「资产管理」/数据库资产.csv 一致，"
        "并确认首页含总览类标题及表头含服务目录编号、服务名称、合计费用/本月核算金额。",
    )
